"""Convert the technical documentation (Markdown) to an editable Word file.

    pip install python-docx
    python submission/tools/md_to_docx.py [source.md [target.docx]]

Headings, paragraphs (bold, italic, code), bullet and numbered lists, tables, code blocks, call-outs and
images are kept. The layout is simpler than the PDF's; the PDF is the version to send.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "docs" / "TECHNICAL_DOCUMENTATION.md"
TARGET = (
    Path(sys.argv[2]).resolve()
    if len(sys.argv) > 2
    else ROOT / "submission" / "StudyRoute_Technical_Documentation.docx"
)
NAVY, TEAL, ZEBRA = "0E2A47", "0B8A86", "F0F5F8"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")


def shade(element, fill: str) -> None:
    props = element.get_or_add_tcPr() if hasattr(element, "get_or_add_tcPr") else element.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def left_bar(paragraph, colour: str) -> None:
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for key, value in (("val", "single"), ("sz", "18"), ("space", "6"), ("color", colour)):
        left.set(qn(f"w:{key}"), value)
    border.append(left)
    paragraph._p.get_or_add_pPr().append(border)


def add_inline(paragraph, text: str, size: float | None = None, colour: str | None = None, bold: bool = False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        run_text, run_bold, run_italic, code = part, bold, False, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run_text, run_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            run_text, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run_text, run_italic = part[1:-1], True
        run = paragraph.add_run(run_text)
        run.bold, run.italic = run_bold, run_italic
        if code:
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)
        if colour:
            run.font.color.rgb = RGBColor.from_string(colour)


def add_table(document, rows: list[list[str]]) -> None:
    columns = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            text = cells[c] if c < len(cells) else ""
            if r == 0:
                add_inline(paragraph, text, size=8.5, colour="FFFFFF", bold=True)
                shade(cell._tc, NAVY)
            else:
                add_inline(paragraph, text, size=8.5)
                if r % 2 == 0:
                    shade(cell._tc, ZEBRA)
    document.add_paragraph()


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_footer(section, label: str) -> None:
    paragraph = section.footer.paragraphs[0]
    run = paragraph.add_run(label + "   |   page ")
    run.font.size = Pt(8)
    field = paragraph.add_run()
    field.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        field._r.append(element)


def main() -> int:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    for name, size, colour in (("Title", 26, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11, NAVY)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.bold = "Arial", Pt(size), True
        style.font.color.rgb = RGBColor.from_string(colour)
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(0.8)
    section.top_margin = section.bottom_margin = Inches(0.8)
    add_footer(section, "SR-TEC-2026-01  |  Final edition  |  Synthetic data, not production assurance")

    i = 0
    seen_title = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
        elif stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(block))
            run.font.name, run.font.size = "Consolas", Pt(8)
            shade(paragraph._p, ZEBRA)
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|[\s:|-]+\|", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(document, rows)
        elif stripped.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            image = (SOURCE.parent / match.group(2)).resolve() if match else None
            if image and image.exists():
                document.add_picture(str(image), width=Inches(5.9))
            i += 1
        elif stripped.startswith(">"):
            parts = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                parts.append(lines[i].strip().lstrip(">").strip())
                i += 1
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(p for p in parts if p), size=9.5)
            left_bar(paragraph, TEAL)
            shade(paragraph._p, "EAF4F4")
        elif re.match(r"#{1,4} ", stripped):
            level = len(stripped.split(" ")[0])
            text = stripped[level + 1:]
            if level == 1 and not seen_title:
                seen_title = True
                document.add_paragraph("STUDYROUTE  /  INTERNATIONAL TRACK  /  TOPIC C").runs[0].font.color.rgb = RGBColor.from_string(TEAL)
                document.add_paragraph("Technical Documentation", style="Title")
                p = document.add_paragraph("Architecture, core algorithms, trained components and security design")
                p.runs[0].font.size = Pt(12)
                document.add_paragraph("SR-TEC-2026-01  |  Final edition  |  18 October 2026  |  Team TeamEkxin")
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            else:
                document.add_heading(text, level=min(level - 1, 3) if level > 1 else 1)
            i += 1
        elif re.match(r"(- |\* )", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            i += 1
        elif re.match(r"\d+\. ", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
        else:
            parts = [stripped]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"(#|\||>|```|!\[|- |\* |\d+\. )", lines[i].strip()):
                parts.append(lines[i].strip())
                i += 1
            add_inline(document.add_paragraph(), " ".join(parts))
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    properties = document.core_properties
    properties.author = "TeamEkxin"
    properties.last_modified_by = "TeamEkxin"
    properties.title = "StudyRoute: technical documentation"
    properties.subject = "FinTechathon 2026, International Track, Topic C"
    document.save(str(TARGET))
    print(f"wrote {TARGET.name}: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables, "
          f"{len(document.inline_shapes)} images")
    return 0


if __name__ == "__main__":
    sys.exit(main())
